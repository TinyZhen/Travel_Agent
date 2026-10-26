"""Build the revised project report docs/Roameo_Project_Report.docx from report_text.py.

The team's original report (docs/Roameo_Project_Report.original.docx, never changed) is the
template: its page size, margins, styles, footer and numbering part are kept, and its body is
replaced. The text lives in report_text.py so it can be read, diffed and checked. Figures come
from report/figures/: the team's images extracted from the original, crops of the app's
screenshots at native pixels, and the deck's vector figures rendered at 300 dpi.

    uv run --no-project --with python-docx==1.1.2 python -I docs/presentation/report/build_report.py

Run from the repository root. Pass --figures-only to refresh the figures without writing the docx.
"""

import copy
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import report_text as T  # noqa: E402

REPO = HERE.parent.parent.parent
ORIGINAL = REPO / "docs/Roameo_Project_Report.original.docx"
OUTPUT = REPO / "docs/Roameo_Project_Report.docx"
FIGURES = HERE / "figures"
SCREENSHOTS = REPO / "docs/screenshots"
SLIDE_FIGURES = REPO / "docs/presentation/slides/figures"
LOGO = REPO / "docs/presentation/slides/logos/Blue-UMB-logo-lockup.png"

# Colours (the deck's calm palette): near-black type, UMB blue for structure, slide-title blue
# for headings, navy for the title, grey for captions.
INK = RGBColor(0x1A, 0x1A, 0x1A)
UMB = RGBColor(0x00, 0x5A, 0x8B)
TITLEBLUE = RGBColor(0x08, 0x45, 0x7E)
NAVY = RGBColor(0x13, 0x22, 0x45)
GREY = RGBColor(0x5B, 0x6B, 0x70)
HEADER_FILL = "DCE6EF"
ROW_FILL = "F5F7FA"

MAX_W, MAX_H, MIN_DPI = 6.5, 6.0, 200  # inches, and the lowest density a screenshot is shown at

# Crops of the app's screenshots, in source pixels (left, top, right, bottom). The originals are
# 2x captures, so every crop keeps its native pixels. "team_imageN" is image N of the original docx.
CROPS = {
    "alex_plan.png": ("team_image5.png", (1056, 219, 2672, 2150)),
    "alex_chat.png": ("team_image5.png", (210, 236, 1005, 650)),
    "alex_steps_research.png": ("team_image5.png", (1056, 5250, 2672, 6120)),
    "alex_steps_critic.png": ("team_image5.png", (1056, 6130, 2672, 6900)),
    "alex_budget.png": ("team_image5.png", (1056, 220, 2672, 790)),
    "alex_question.png": ("team_image5.png", (210, 1375, 1010, 1965)),
    "guide_crop.png": ("team_image3.png", (0, 0, 1616, 1800)),
}
STACKS = {  # two crops of one screenshot stacked into one figure, still at native pixels
    "detroit_crop.png": ("team_image4.png", [(210, 240, 1005, 545), (210, 1125, 1005, 1385)]),
}
VECTOR = ["fig_ideas", "fig_workflow", "fig_price", "fig_results", "fig_dataflow"]


# ---------------------------------------------------------------------------- figures

def prepare_figures() -> list[str]:
    """Extract the team's images, cut the crops and render the vector figures. Returns the
    names of the vector figures that are still missing."""
    FIGURES.mkdir(exist_ok=True)
    with zipfile.ZipFile(ORIGINAL) as archive:
        for n in range(1, 6):
            target = FIGURES / f"team_image{n}.png"
            if not target.exists():
                target.write_bytes(archive.read(f"word/media/image{n}.png"))
    for name, (source, box) in CROPS.items():
        Image.open(FIGURES / source).crop(box).save(FIGURES / name, optimize=True)
    for name, (source, boxes) in STACKS.items():
        src = Image.open(FIGURES / source)
        parts = [src.crop(b) for b in boxes]
        gap = 24
        out = Image.new("RGB", (max(p.width for p in parts), sum(p.height for p in parts) + gap), "white")
        y = 0
        for p in parts:
            out.paste(p, (0, y))
            y += p.height + gap
        out.save(FIGURES / name, optimize=True)
    missing = []
    for stem in VECTOR:
        pdf = SLIDE_FIGURES / f"{stem}.pdf"
        png = FIGURES / f"{stem}.png"
        if not pdf.exists():
            missing.append(stem)
            continue
        if not png.exists() or png.stat().st_mtime < pdf.stat().st_mtime:
            subprocess.run(["pdftoppm", "-r", "300", "-png", "-singlefile", str(pdf), str(png.with_suffix(""))],
                           check=True)
    return missing


def picture_size(path: Path, max_w: float = MAX_W, max_h: float = MAX_H) -> tuple[float, float]:
    """Width and height in inches: as large as the box allows, but never below MIN_DPI."""
    w, h = Image.open(path).size
    dpi = max(w / max_w, h / max_h, MIN_DPI)
    return w / dpi, h / dpi


# ---------------------------------------------------------------------------- styles

def set_style_font(style, size=None, bold=None, italic=None, color=None, font=None):
    f = style.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color
    if font is not None:
        f.name = font
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), font)


def add_numbering(doc) -> tuple[int, int]:
    """Add one bullet and one decimal list definition. Returns their abstractNum ids."""
    numbering = doc.part.numbering_part.element
    ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    bullet_id, decimal_id = max(ids, default=-1) + 1, max(ids, default=-1) + 2
    bullet = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{bullet_id}"><w:multiLevelType w:val="hybridMultilevel"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="•"/>'
        '<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="360" w:hanging="270"/></w:pPr>'
        '<w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:color w:val="005A8B"/></w:rPr></w:lvl></w:abstractNum>')
    decimal = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{decimal_id}"><w:multiLevelType w:val="hybridMultilevel"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/>'
        '<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="360" w:hanging="300"/></w:pPr>'
        '<w:rPr><w:b/><w:color w:val="005A8B"/></w:rPr></w:lvl></w:abstractNum>')
    first_num = numbering.find(qn("w:num"))
    for el in (bullet, decimal):
        if first_num is not None:
            first_num.addprevious(el)
        else:
            numbering.append(el)
    return bullet_id, decimal_id


def new_num(doc, abstract_id: int, restart: bool) -> int:
    numbering = doc.part.numbering_part.element
    num_id = max([int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))], default=0) + 1
    override = ('<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride>' if restart else "")
    numbering.append(parse_xml(
        f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{abstract_id}"/>{override}</w:num>'))
    return num_id


def set_num(paragraph_or_style_el, num_id: int):
    ppr = paragraph_or_style_el.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is None:
        numpr = parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr>')
        ppr.insert(0, numpr)
    else:
        numpr.find(qn("w:numId")).set(qn("w:val"), str(num_id))


def setup_styles(doc):
    styles = doc.styles
    # Default type: near-black Calibri 11.
    rpr_default = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    rpr_default.find(qn("w:color")).set(qn("w:val"), "1A1A1A")
    normal = styles["Normal"]
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    set_style_font(styles["Heading 1"], size=17, bold=True, color=TITLEBLUE)
    styles["Heading 1"].paragraph_format.keep_with_next = True
    set_style_font(styles["Heading 2"], size=13, bold=True, color=INK)
    styles["Heading 2"].paragraph_format.keep_with_next = True
    h3 = styles["Heading 3"]
    set_style_font(h3, size=11.5, bold=True, color=UMB)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(3)
    h3.paragraph_format.keep_with_next = True
    h3.element.find(qn("w:semiHidden")) is not None and h3.element.remove(h3.element.find(qn("w:semiHidden")))

    bullet_id, decimal_id = add_numbering(doc)
    for name, abstract in (("List Bullet", bullet_id), ("List Number", decimal_id)):
        st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = styles["Normal"]
        st.paragraph_format.space_after = Pt(3)
        set_num(st.element, new_num(doc, abstract, restart=False))
    cap = styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    cap.base_style = styles["Normal"]
    set_style_font(cap, size=9, italic=False, color=GREY)
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(12)
    code = styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    code.base_style = styles["Normal"]
    set_style_font(code, size=9, font="Consolas", color=INK)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.left_indent = Inches(0.2)
    code.element.get_or_add_pPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="F2F4F7"/>'))
    inline = styles.add_style("Inline Code", WD_STYLE_TYPE.CHARACTER)
    set_style_font(inline, size=9.5, font="Consolas")
    tt = styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    tt.base_style = styles["Normal"]
    set_style_font(tt, size=9.5)
    tt.paragraph_format.space_after = Pt(2)
    tt.paragraph_format.line_spacing = 1.0
    return decimal_id


# ---------------------------------------------------------------------------- inline text

TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\))")


def add_hyperlink(paragraph, text: str, url: str, bold=False):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = parse_xml(f'<w:hyperlink {nsdecls("w", "r")} r:id="{r_id}" w:history="1"/>')
    run = parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:rStyle w:val="Hyperlink"/>{"<w:b/>" if bold else ""}</w:rPr></w:r>')
    t = parse_xml(f'<w:t {nsdecls("w")} xml:space="preserve"/>')
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_inline(paragraph, text: str, bold=False, size=None, color=None):
    """Write text with **bold**, `code` (Consolas, exempt from the writing check) and
    [shown text](link target)."""
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1], style="Inline Code")
        elif part.startswith("[") and "](" in part:
            shown, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part).groups()
            add_hyperlink(paragraph, shown, url, bold=bold)
            continue
        else:
            run = paragraph.add_run(part)
            if bold:
                run.bold = True
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color


# ---------------------------------------------------------------------------- blocks

class Builder:
    def __init__(self, doc, decimal_id):
        self.doc = doc
        self.decimal_id = decimal_id
        self.figure_no = 0
        self.table_no = 0
        self.headings: list[tuple[int, str]] = []
        self.missing: list[str] = []

    def para(self, text, style=None, align=None, space_after=None, keep=False):
        p = self.doc.add_paragraph(style=style)
        add_inline(p, text)
        if align:
            p.alignment = align
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        if keep:
            p.paragraph_format.keep_with_next = True
        return p

    def heading(self, level, text):
        self.headings.append((level, text))
        return self.doc.add_paragraph(text, style=f"Heading {level}")

    def bullets(self, items):
        for item in items:
            self.para(item, style="List Bullet")

    def numbers(self, items):
        num_id = new_num(self.doc, self.decimal_id, restart=True)
        for item in items:
            p = self.para(item, style="List Number")
            set_num(p._p, num_id)

    def code(self, lines):
        for i, line in enumerate(lines):
            p = self.doc.add_paragraph(style="Code")
            p.add_run(line)
            if i == len(lines) - 1:
                p.paragraph_format.space_after = Pt(8)

    def caption(self, kind, text, keep=False):
        if kind == "Figure":
            self.figure_no += 1
            label = f"Figure {self.figure_no}."
        else:
            self.table_no += 1
            label = f"Table {self.table_no}."
        p = self.doc.add_paragraph(style="Caption")
        r = p.add_run(label + " ")
        r.bold = True
        r.font.color.rgb = UMB
        add_inline(p, text)
        if keep:
            p.paragraph_format.keep_with_next = True
        return p

    def figure(self, name, caption, max_w=MAX_W, max_h=MAX_H):
        path = FIGURES / name
        if not path.exists():
            self.missing.append(name)
            return
        w, h = picture_size(path, max_w, max_h)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(str(path), width=Inches(w), height=Inches(h))
        cp = self.caption("Figure", caption)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def table(self, widths, header, rows, caption, bold_first_col=False):
        self.caption("Table", caption, keep=True)
        table = self.doc.add_table(rows=1 + len(rows), cols=len(widths))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        tbl = table._tbl
        tblpr = tbl.tblPr
        total = sum(widths)
        tblpr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="{int(total * 1440)}" w:type="dxa"/>'))
        tblpr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
        tblpr.append(parse_xml(
            f'<w:tblBorders {nsdecls("w")}><w:top w:val="nil"/><w:left w:val="nil"/><w:bottom w:val="nil"/>'
            '<w:right w:val="nil"/><w:insideH w:val="nil"/><w:insideV w:val="nil"/></w:tblBorders>'))
        tblpr.append(parse_xml(
            f'<w:tblCellMar {nsdecls("w")}><w:top w:w="50" w:type="dxa"/><w:left w:w="90" w:type="dxa"/>'
            '<w:bottom w:w="50" w:type="dxa"/><w:right w:w="90" w:type="dxa"/></w:tblCellMar>'))
        grid = tbl.tblGrid
        for col, w in zip(grid.findall(qn("w:gridCol")), widths):
            col.set(qn("w:w"), str(int(w * 1440)))
        for r_i, values in enumerate([header] + rows):
            row = table.rows[r_i]
            if r_i == 0:
                row._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            for c_i, (cell, value) in enumerate(zip(row.cells, values)):
                cell.width = Inches(widths[c_i])
                tcpr = cell._tc.get_or_add_tcPr()
                fill = HEADER_FILL if r_i == 0 else (ROW_FILL if r_i % 2 == 0 else None)
                if fill:
                    tcpr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>'))
                cell.paragraphs[0].style = self.doc.styles["Table Text"]
                lines = value if isinstance(value, list) else [value]
                for l_i, line in enumerate(lines):
                    p = cell.paragraphs[0] if l_i == 0 else cell.add_paragraph(style="Table Text")
                    add_inline(p, line, bold=(r_i == 0 or (bold_first_col and c_i == 0)))
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)

    def page_break(self):
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)

    # -- special pages
    def title_page(self, t):
        d = self.doc
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(36)
        w, h = Image.open(LOGO).size
        p.add_run().add_picture(str(LOGO), width=Inches(2.2), height=Inches(2.2 * h / w))
        for text, size, bold, color, after in t["lines"]:
            q = d.add_paragraph()
            q.paragraph_format.space_after = Pt(after)
            add_inline(q, text, bold=bold, size=size, color=color)

    def contents(self):
        """A TOC field whose cached result lists the headings, refreshed by Word on opening."""
        d = self.doc
        p = d.add_paragraph()
        r = p.add_run("Contents")
        r.bold = True
        r.font.size = Pt(14)
        r.font.color.rgb = TITLEBLUE
        self._toc_anchor = d.add_paragraph()

    def fill_contents(self):
        anchor = self._toc_anchor
        entries = [(lvl, txt) for lvl, txt in self.headings if lvl <= 2 and txt not in T.NOT_IN_CONTENTS]
        paras = []
        for i, (lvl, txt) in enumerate(entries):
            style = "TOC1" if lvl == 1 else "TOC2"
            begin = ('<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
                     '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-2" \\h \\z \\u </w:instrText></w:r>'
                     '<w:r><w:fldChar w:fldCharType="separate"/></w:r>') if i == 0 else ""
            end = '<w:r><w:fldChar w:fldCharType="end"/></w:r>' if i == len(entries) - 1 else ""
            safe = txt.replace("&", "&amp;").replace("<", "&lt;")
            paras.append(parse_xml(
                f'<w:p {nsdecls("w")}><w:pPr><w:pStyle w:val="{style}"/></w:pPr>{begin}'
                f'<w:r><w:t xml:space="preserve">{safe}</w:t></w:r>{end}</w:p>'))
        for el in paras:
            anchor._p.addprevious(el)
        anchor._p.getparent().remove(anchor._p)


# ---------------------------------------------------------------------------- assemble

def clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def drop_unused_rels(doc):
    xml = doc.element.xml
    for r_id, rel in list(doc.part.rels.items()):
        if rel.reltype in (RT.IMAGE, RT.HYPERLINK) and f'"{r_id}"' not in xml:
            del doc.part.rels[r_id]


def set_footer(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    def run_xml(inner):
        return parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:color w:val="5B6B70"/><w:sz w:val="16"/>'
                         f'<w:szCs w:val="16"/></w:rPr>{inner}</w:r>')
    p._p.append(run_xml('<w:t xml:space="preserve">Roameo, PEAAII AI Applications Hackathon 2026, page </w:t>'))
    p._p.append(run_xml('<w:fldChar w:fldCharType="begin"/>'))
    p._p.append(run_xml('<w:instrText xml:space="preserve"> PAGE </w:instrText>'))
    p._p.append(run_xml('<w:fldChar w:fldCharType="separate"/>'))
    p._p.append(run_xml('<w:t>1</w:t>'))
    p._p.append(run_xml('<w:fldChar w:fldCharType="end"/>'))


def set_update_fields(doc):
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        settings.append(parse_xml(f'<w:updateFields {nsdecls("w")} w:val="true"/>'))


def render(builder: Builder):
    for block in T.BLOCKS:
        kind, *args = block
        if kind == "title":
            builder.title_page(args[0])
        elif kind == "contents":
            builder.contents()
        elif kind == "pagebreak":
            builder.page_break()
        elif kind in ("h1", "h2", "h3"):
            builder.heading(int(kind[1]), args[0])
        elif kind == "p":
            builder.para(args[0])
        elif kind == "bullets":
            builder.bullets(args[0])
        elif kind == "numbers":
            builder.numbers(args[0])
        elif kind == "code":
            builder.code(args[0])
        elif kind == "figure":
            builder.figure(*args)
        elif kind == "table":
            builder.table(**args[0])
        else:
            raise ValueError(kind)
    builder.fill_contents()


def main() -> int:
    missing = prepare_figures()
    if missing:
        print("vector figures not found yet:", ", ".join(missing))
    if "--figures-only" in sys.argv:
        return 0
    doc = Document(str(ORIGINAL))
    clear_body(doc)
    decimal_id = setup_styles(doc)
    builder = Builder(doc, decimal_id)
    render(builder)
    set_footer(doc)
    set_update_fields(doc)
    drop_unused_rels(doc)
    cp = doc.core_properties
    cp.title = T.DOC_TITLE
    cp.subject = T.DOC_SUBJECT
    cp.author = "Zhen Lu, Noctis Zhang"
    cp.last_modified_by = "Zhen Lu, Noctis Zhang"
    cp.keywords = "Roameo, agentic AI, trip planning"
    cp.comments = ""
    doc.save(str(OUTPUT))
    words = sum(len(re.findall(r"[A-Za-z0-9$][\w'$.,/%-]*", p.text)) for p in Document(str(OUTPUT)).paragraphs)
    table_words = sum(len(re.findall(r"[A-Za-z0-9$][\w'$.,/%-]*", c.text))
                      for t in Document(str(OUTPUT)).tables for r in t.rows for c in r.cells)
    print(f"wrote {OUTPUT.relative_to(REPO)}: {builder.figure_no} figures, {builder.table_no} tables, "
          f"{words} words in paragraphs, {table_words} in tables")
    if builder.missing:
        print("figures skipped (file missing):", ", ".join(builder.missing))
    return 0


if __name__ == "__main__":
    sys.exit(main())
